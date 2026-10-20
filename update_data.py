#!/usr/bin/env python3
"""
Auto Data Feed for the College Football Predictive Betting Model
================================================================

Pulls current-season numbers from the free CollegeFootballData API and writes
them into CFB_Betting_Model.xlsx:

    Weekly Input tab :  SP+ Overall / Offense / Defense, Yards-per-play
    Auto Feed tab    :  Talent, Recent Form (4wk), Success Rate,
                        Explosiveness, Offense PPA, Defense PPA,
                        Def Havoc, Off Havoc Allowed
    Schedule tab     :  every FBS game with date, site and result
    New tab          :  the per-team stat sheet (27 columns)

HOW A FEW "New" TAB COLUMNS ARE WORKED OUT
    Talent Points       the model's own talent weighting, read live from
                        Power Ratings (so it follows Weights & Params)
    Turnover margin     takeaways minus giveaways, season to date
    Time of possession  average per game, MM:SS
    FG / Punt yards     added up from the individual kickers' season stats

SETUP (once)
    1.  pip install requests openpyxl
    2.  Get a FREE key at https://collegefootballdata.com/key  (emailed instantly)
    3.  Paste it into API_KEY below. The season year is set automatically to the
        current calendar year; WORKBOOK only needs changing if you rename the file.

RUN
    Normally you never run this by hand: double-click the "CFB Betting Model" desktop
    shortcut (or "Open CFB Model.bat"). It runs this updater, then opens the workbook.

    The updater only calls the API once per day, to save API credits. Later opens
    that day skip straight to the workbook. To force a refresh anyway, close the
    workbook and run:   python update_data.py --force

The model recalculates automatically the next time you open the workbook.
Any team whose API name doesn't match the sheet is printed at the end so you
can add it to NAME_FIXES.
"""

import argparse
import os
import re
import sys
import time
from copy import copy
from datetime import date, datetime

try:                                    # US Eastern for kickoff times (CFB convention)
    from zoneinfo import ZoneInfo
    _ET = ZoneInfo("America/New_York")
except Exception:
    _ET = None

API_KEY  = "MUUUNFo1ICHMRpyBwImFWe6SKjJMMST6Zd9gmP2QHf/Wc+ykaj54K/0F6ECZNXfL"     # <-- from collegefootballdata.com/key
YEAR     = datetime.now().year            # always the current calendar year
WORKBOOK = "CFB_Betting_Model.xlsx"        # <-- path to the workbook

HERE = os.path.dirname(os.path.abspath(__file__))
WORKBOOK = os.path.join(HERE, WORKBOOK)
STAMP = os.path.join(HERE, ".last_update")   # date of the last successful update

BASE = "https://api.collegefootballdata.com"

# API "school" name  ->  workbook team name (only the ones that differ)
NAME_FIXES = {
    "Connecticut": "UConn",
    "Louisiana Monroe": "Louisiana-Monroe",
    "UL Monroe": "Louisiana-Monroe",
    "San José State": "San Jose State",
    "Hawai'i": "Hawaii",
    "Southern Mississippi": "Southern Miss",
    "App State": "Appalachian State",
    "Massachusetts": "UMass",
    "Sam Houston State": "Sam Houston",
    "Louisiana Lafayette": "Louisiana",
}


def _need(mod):
    try:
        return __import__(mod)
    except ImportError:
        sys.exit(f"Missing package '{mod}'. Run:  pip install requests openpyxl")


requests = _need("requests")
openpyxl = _need("openpyxl")
from openpyxl.formula.translate import Translator


FAILED = []     # endpoints that never answered - reported loudly at the end


def get(path, **params):
    """GET a CFBD endpoint and return parsed JSON (empty list on failure)."""
    headers = {"Authorization": f"Bearer {API_KEY}", "Accept": "application/json"}
    for attempt in range(3):
        try:
            r = requests.get(f"{BASE}{path}", headers=headers, params=params, timeout=60)
            if r.status_code == 401:
                sys.exit("401 Unauthorized — check your API_KEY.")
            if r.status_code == 429:
                sys.exit("429 Too Many Requests — the free API key's monthly call limit is used up.")
            r.raise_for_status()
            return r.json()
        except requests.RequestException as e:
            if attempt == 2:
                print(f"  ! {path} failed: {e}")
                FAILED.append(path)
                return []
            time.sleep(5)
    return []


def check_api():
    """Fail fast (instead of 10+ minutes of timeouts) when CFBD itself is down."""
    try:
        requests.get(f"{BASE}/conferences", timeout=30,
                     headers={"Authorization": f"Bearer {API_KEY}"})
    except requests.RequestException:
        sys.exit(
            "\nCollegeFootballData.com is not responding right now (their server is down,\n"
            "not your computer). Nothing was changed - try again in a few hours.\n"
            "Status / announcements: https://collegefootballdata.com"
        )


# Every team-list range in the workbook ends on the last team row (e.g. $B$2:$B$137).
TEAM_RANGE_END = re.compile(r"(:\$?[A-Z]{1,3}\$)(\d+)\b")


def add_new_fbs_teams(wb, fbs_conf):
    """Append FBS teams the API knows about but the sheet doesn't (e.g. programs that
    just moved up), copy the formulas down, and stretch every team range to cover them."""
    wi, af, pr = wb["Weekly Input"], wb["Auto Feed"], wb["Power Ratings"]
    have = {wi.cell(r, 1).value for r in range(2, wi.max_row + 1)}
    new = sorted(t for t in fbs_conf if t and t not in have)
    if not new:
        return []
    old_last = pr.max_row
    for i, team in enumerate(new, start=1):
        r = old_last + i
        wi.cell(r, 1, team)
        af.cell(r, 1, team); af.cell(r, 2, fbs_conf[team])
        pr.cell(r, 2, team)
        for ws in (wi, af, pr):
            for c in range(1, ws.max_column + 1):
                src, dst = ws.cell(old_last, c), ws.cell(r, c)
                if src.has_style:
                    dst._style = copy(src._style)
                if isinstance(src.value, str) and src.value.startswith("="):
                    dst.value = Translator(src.value, origin=src.coordinate).translate_formula(dst.coordinate)
    new_last = old_last + len(new)

    def stretch(f):
        return TEAM_RANGE_END.sub(
            lambda m: f"{m.group(1)}{new_last}" if int(m.group(2)) == old_last else m.group(0), f)

    for ws in wb:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    c.value = stretch(c.value)
    for name in wb.defined_names.values():
        name.attr_text = stretch(name.attr_text)
    return new


def last_update():
    try:
        with open(STAMP) as f:
            return f.read().strip()
    except OSError:
        return ""


def workbook_locked():
    """True if Excel (or OneDrive) has the workbook open, so it can't be saved."""
    try:
        with open(WORKBOOK, "r+b"):
            return False
    except PermissionError:
        return True
    except FileNotFoundError:
        sys.exit(f"Can't find {WORKBOOK}.")


def norm(name):
    return NAME_FIXES.get(name, name)


# ============================================================================
#  "New" tab - the extra per-team stat sheet
# ============================================================================

NEW_SHEET = "New"

# Column headings, in order, starting in column A of row 2.
NEW_HEADERS = [
    "Team", "Games played", "Talent Composite", "Talent Points", "ESPN FPI",
    "Overall SP+", "Turnover margin",
    "Off SP+", "Off PPA", "Off YPP", "Off Havoc Rate", "Off Success Rate",
    "First Downs", "Time of possession",
    "Def SP+", "Def PPA", "Def YPP", "Def Havoc Rate", "Def Success Rate",
    "Opponent First Downs",
    "Special Teams SP+", "FG (Made/Attempted)", "Avg Punt Yards",
    "Avg Punt Return Yds gained", "Avg Kick Return Yds gained",
    "opponent Average Punt Return Yds", "opponent average kick return Yds",
]


def player_totals(category, stat_types):
    """Sum a player-stat category to team level: {team: {statType: total}}."""
    out = {}
    for row in get("/stats/player/season", year=YEAR, category=category):
        st = row.get("statType")
        if st not in stat_types:
            continue
        try:
            val = float(row.get("stat"))
        except (TypeError, ValueError):
            continue
        out.setdefault(norm(row.get("team", "")), {}).setdefault(st, 0.0)
        out[norm(row.get("team", ""))][st] += val
    return out


def _ratio(num, den, nd=2):
    try:
        return round(num / den, nd) if den else None
    except TypeError:
        return None


def _mmss(seconds):
    """Seconds -> "MM:SS" for the time-of-possession column."""
    if not seconds:
        return None
    s = int(round(seconds))
    return f"{s // 60}:{s % 60:02d}"


def new_sheet_rows(teams, data, season, kicking, punting):
    """Build one output row per team, in the order of NEW_HEADERS."""
    rows = []
    for team in teams:
        s = season.get(team, {})
        d = data.get(team, {})
        k = kicking.get(team, {})
        pu = punting.get(team, {})
        g = s.get("games") or 0

        def_plays = (s.get("rushingAttemptsOpponent") or 0) + (s.get("passAttemptsOpponent") or 0)
        to_margin = None
        if s.get("turnovers") is not None and s.get("turnoversOpponent") is not None:
            to_margin = s["turnoversOpponent"] - s["turnovers"]

        rows.append([
            team,
            g or None,
            d.get("talent"),
            None,                                   # Talent Points - formula, filled below
            d.get("fpi"),
            d.get("sp_overall"),
            to_margin,
            d.get("sp_off"),
            d.get("off_ppa"),
            d.get("ypp"),
            d.get("off_havoc"),
            d.get("success"),
            s.get("firstDowns"),
            _mmss(_ratio(s.get("possessionTime"), g, 0)),
            d.get("sp_def"),
            d.get("def_ppa"),
            _ratio(s.get("totalYardsOpponent"), def_plays),
            d.get("def_havoc"),
            d.get("def_success"),
            s.get("firstDownsOpponent"),
            d.get("sp_st"),
            f"{int(k.get('FGM', 0))}/{int(k.get('FGA', 0))}" if k.get("FGA") else None,
            _ratio(pu.get("YDS"), pu.get("NO"), 1),
            _ratio(s.get("puntReturnYards"), s.get("puntReturns"), 1),
            _ratio(s.get("kickReturnYards"), s.get("kickReturns"), 1),
            _ratio(s.get("puntReturnYardsOpponent"), s.get("puntReturnsOpponent"), 1),
            _ratio(s.get("kickReturnYardsOpponent"), s.get("kickReturnsOpponent"), 1),
        ])
        assert len(rows[-1]) == len(NEW_HEADERS)
    return rows


def write_new_sheet(wb, rows, pr_last_row):
    """Create/refresh the 'New' tab: headings on row 2, one team per row from row 3."""
    from openpyxl.styles import Alignment, Font

    if NEW_SHEET in wb.sheetnames:
        ws = wb[NEW_SHEET]
    else:
        ws = wb.create_sheet(NEW_SHEET, 0)
        ws.freeze_panes = "B3"
        ws.column_dimensions["A"].width = 22
    wb.move_sheet(ws, -wb.sheetnames.index(NEW_SHEET))   # keep it as the first tab

    wide = max(ws.max_column, len(NEW_HEADERS))
    for r in range(2, ws.max_row + 1):                   # clear the old table
        for col in range(1, wide + 1):
            ws.cell(r, col).value = None

    for col, head in enumerate(NEW_HEADERS, start=1):
        c = ws.cell(2, col, head)
        c.font = Font(bold=True)
        c.alignment = Alignment(wrap_text=True, vertical="bottom")

    for i, row in enumerate(rows, start=3):
        for col, val in enumerate(row, start=1):
            ws.cell(i, col, round(val, 3) if isinstance(val, float) else val)
        # Talent Points comes from the model's own weighting on Power Ratings.
        ws.cell(i, 4, f"=IFERROR(INDEX('Power Ratings'!$L$2:$L${pr_last_row},"
                      f"MATCH($A{i},'Power Ratings'!$B$2:$B${pr_last_row},0)),\"\")")
    return ws


def main():
    ap = argparse.ArgumentParser(description="Pull the latest CFB data into the workbook.")
    ap.add_argument("--force", action="store_true",
                    help="update even if it already ran today (uses API credits)")
    args = ap.parse_args()

    if API_KEY == "PASTE_YOUR_FREE_KEY_HERE":
        sys.exit("Set API_KEY first (free key at https://collegefootballdata.com/key).")

    today = date.today().isoformat()
    if last_update() == today and not args.force:
        print(f"Data already updated today ({today}) - skipping to save API credits.")
        return

    # Check before spending any API calls: a locked workbook can't be saved.
    if workbook_locked():
        print("The workbook is already open, so it can't be updated right now.\n"
              "Close it everywhere (Excel AND any browser tab) and open it again\n"
              "to pull today's data.")
        sys.exit(2)

    print(f"Pulling {YEAR} data from CollegeFootballData ...")
    check_api()

    # per-team accumulator
    data = {}   # team -> dict of metrics

    def slot(team):
        return data.setdefault(team, {})

    # ---- SP+ ratings ----
    for row in get("/ratings/sp", year=YEAR):
        t = row.get("team")
        if not t:
            continue
        d = slot(norm(t))
        d["sp_overall"] = row.get("rating")
        off = row.get("offense") or {}
        dfn = row.get("defense") or {}
        d["sp_off"] = off.get("rating")
        d["sp_def"] = dfn.get("rating")
        d["sp_st"] = (row.get("specialTeams") or {}).get("rating")

    # ---- ESPN FPI ----
    for row in get("/ratings/fpi", year=YEAR):
        slot(norm(row.get("team", "")))["fpi"] = row.get("fpi")

    # ---- Talent composite ----
    for row in get("/talent", year=YEAR):
        d = slot(norm(row.get("team", "")))
        d["talent"] = row.get("talent")

    # ---- PPA (predicted points added) ----
    for row in get("/ppa/teams", year=YEAR):
        d = slot(norm(row.get("team", "")))
        off = (row.get("offense") or {})
        dfn = (row.get("defense") or {})
        d["off_ppa"] = off.get("overall")
        d["def_ppa"] = dfn.get("overall")

    # ---- Advanced season stats: success rate + explosiveness + havoc ----
    for row in get("/stats/season/advanced", year=YEAR):
        d = slot(norm(row.get("team", "")))
        off = (row.get("offense") or {})
        dfn = (row.get("defense") or {})
        d["success"] = off.get("successRate")
        d["def_success"] = dfn.get("successRate")
        d["explode"] = off.get("explosiveness")
        d["def_havoc"] = (dfn.get("havoc") or {}).get("total")   # defense creates havoc (higher better)
        d["off_havoc"] = (off.get("havoc") or {}).get("total")   # offense allows havoc (lower better)

    # ---- Season stats (also feeds most of the "New" tab) ----
    season = {}     # team -> {statName: value}
    for row in get("/stats/season", year=YEAR):
        t = norm(row.get("team", ""))
        if row.get("statName") is not None:
            season.setdefault(t, {})[row["statName"]] = row.get("statValue")

    # yards per play = totalYards / (rush att + pass att)
    for t, s in season.items():
        plays = (s.get("rushingAttempts") or 0) + (s.get("passAttempts") or 0)
        if plays:
            slot(t)["ypp"] = round((s.get("totalYards") or 0) / plays, 2)

    # ---- Recent form: avg scoring margin over last <=4 completed games ----
    games = get("/games", year=YEAR, seasonType="regular")
    hist = {}   # team -> list of (week, margin)
    fbs_conf = {}   # every team the API classes as FBS this season -> conference
    for g in games:
        for side in ("home", "away"):
            if g.get(f"{side}Classification") == "fbs":
                fbs_conf[norm(g.get(f"{side}Team", ""))] = g.get(f"{side}Conference") or ""
        if not g.get("completed"):
            continue
        hp, ap = g.get("homePoints"), g.get("awayPoints")
        if hp is None or ap is None:
            continue
        h, a, wk = norm(g.get("homeTeam", "")), norm(g.get("awayTeam", "")), g.get("week", 0)
        hist.setdefault(h, []).append((wk, hp - ap))
        hist.setdefault(a, []).append((wk, ap - hp))
    for t, lst in hist.items():
        lst.sort(key=lambda x: x[0])
        last4 = [m for _, m in lst[-4:]]
        if last4:
            slot(t)["form"] = round(sum(last4) / len(last4), 1)

    # ---- kicking and punting totals for the "New" tab ----
    kicking = player_totals("kicking", {"FGM", "FGA"})
    punting = player_totals("punting", {"YDS", "NO"})

    # ---- write into workbook ----
    print(f"Opening {WORKBOOK} ...")
    wb = openpyxl.load_workbook(WORKBOOK)
    wi, af = wb["Weekly Input"], wb["Auto Feed"]

    added = add_new_fbs_teams(wb, fbs_conf)
    if added:
        print(f"Added {len(added)} new FBS team(s) to the model: {', '.join(added)}")

    def sheet_index(ws):
        return {ws.cell(r, 1).value: r for r in range(2, ws.max_row + 1)}

    wi_rows, af_rows = sheet_index(wi), sheet_index(af)
    sheet_teams = set(wi_rows) | set(af_rows)

    written, matched, sp_changed = 0, set(), 0
    for team, d in data.items():
        if team in wi_rows:
            r = wi_rows[team]
            old_sp = wi.cell(r, 2).value
            if d.get("sp_overall") is not None and old_sp != round(float(d["sp_overall"]), 3):
                sp_changed += 1
            for col, key in [(2, "sp_overall"), (3, "sp_off"), (4, "sp_def"), (5, "ypp")]:
                if d.get(key) is not None:
                    wi.cell(r, col, round(float(d[key]), 3)); written += 1
            matched.add(team)
        if team in af_rows:
            r = af_rows[team]
            for col, key in [(3, "talent"), (4, "form"), (5, "success"),
                             (6, "explode"), (7, "off_ppa"), (8, "def_ppa"),
                             (9, "def_havoc"), (10, "off_havoc")]:
                if d.get(key) is not None:
                    af.cell(r, col, round(float(d[key]), 4)); written += 1
            matched.add(team)

    # ---- "New" tab: the full per-team stat sheet ----
    teams = [t for t in sorted(wi_rows) if t]
    rows = new_sheet_rows(teams, data, season, kicking, punting)
    write_new_sheet(wb, rows, wb["Power Ratings"].max_row)
    if "Dashboard" in wb.sheetnames:            # still open on the Dashboard
        wb.active = wb.sheetnames.index("Dashboard")
    filled = sum(1 for r in rows if any(v is not None for v in r[1:]))
    print(f"New tab: wrote {len(NEW_HEADERS)} stats for {filled} teams.")

    # ---- Schedule tab: dates for every FBS game (regular + postseason) ----
    post = get("/games", year=YEAR, seasonType="postseason")
    sched = []
    for g in list(games) + list(post):
        if g.get("homeClassification") != "fbs" and g.get("awayClassification") != "fbs":
            continue
        sd = g.get("startDate") or ""
        when = sd[:10]
        if sd:
            try:
                dt = datetime.fromisoformat(sd.replace("Z", "+00:00"))
                dt = dt.astimezone(_ET) if _ET else dt.astimezone()
                day = dt.strftime("%a %b %d")
                if g.get("startTimeTBD"):
                    tm = "TBD"
                else:
                    tm = dt.strftime("%I:%M %p").lstrip("0") + (" ET" if _ET else "")
                when = f"{day}  {tm}"
            except Exception:
                pass
        home, away = norm(g.get("homeTeam", "") or ""), norm(g.get("awayTeam", "") or "")
        wk = "Bowl" if g.get("seasonType") == "postseason" else g.get("week")
        site = "Neutral site" if g.get("neutralSite") else (g.get("venue") or "")
        hp, ap = g.get("homePoints"), g.get("awayPoints")
        result = f"{ap}-{hp}" if (g.get("completed") and hp is not None and ap is not None) else ""
        key = "|".join(sorted([away, home]))
        sched.append((sd, when, wk, away, home, site, result, key))
    sched.sort(key=lambda r: r[0])

    if "Schedule" not in wb.sheetnames:
        s0 = wb.create_sheet("Schedule")
        for col, name in enumerate(["Date", "Week", "Away", "Home", "Site", "Result", "Key"], start=1):
            s0.cell(1, col, name)
    ss = wb["Schedule"]
    if "/games" in FAILED:                      # don't wipe the schedule on a failed download
        print("Schedule: kept last run's games (the /games download failed).")
    else:
        for r in range(2, ss.max_row + 1):      # clear previous rows
            for c in range(1, 8):
                ss.cell(r, c).value = None
        for i, (_, when, wk, away, home, site, result, key) in enumerate(sched, start=2):
            ss.cell(i, 1, when); ss.cell(i, 2, wk); ss.cell(i, 3, away)
            ss.cell(i, 4, home); ss.cell(i, 5, site); ss.cell(i, 6, result); ss.cell(i, 7, key)
        print(f"Schedule: wrote {len(sched)} FBS games.")

    # keep Excel recomputing on open
    from openpyxl.workbook.properties import CalcProperties
    wb.calculation = CalcProperties(calcId=124519, fullCalcOnLoad=True, calcMode="auto")
    try:
        wb.save(WORKBOOK)
    except PermissionError:
        sys.exit(
            f"\nERROR: '{WORKBOOK}' is open or locked, so it can't be updated.\n"
            "  1. Close the workbook EVERYWHERE (browser tab AND any Excel window).\n"
            "  2. Wait a few seconds for OneDrive to finish syncing.\n"
            "  3. Run this again.\n"
            "(Your data was NOT changed - nothing was written.)"
        )

    if not FAILED:      # a partial update is retried on the next open
        with open(STAMP, "w") as f:
            f.write(today)

    print(f"\nDone. Updated {len(matched)} FBS teams ({written} values).")
    unmatched_api = sorted(t for t in data if t not in sheet_teams)
    missing_sheet = sorted(t for t in sheet_teams if t not in matched)
    if unmatched_api:
        print(f"(Skipped {len(unmatched_api)} non-FBS teams the API returns - this is normal.)")
    if missing_sheet:
        # The only warning that matters: an FBS team on the sheet that got no data.
        print("\nWARNING - these FBS teams got NO data (name may not match the API):")
        for t in missing_sheet:
            print("   -", t)
        print("Add a fix to NAME_FIXES near the top of this file if a real team is listed.")
    else:
        print("Every FBS team on the sheet updated cleanly.")
    if "/ratings/sp" not in FAILED and sp_changed == 0:
        print("\nNOTE - SP+ numbers are the same as last time. CollegeFootballData usually posts\n"
              "the new week's SP+ a day or two after ESPN does; run this again later.")
    if FAILED:
        print("\n*** WARNING - these data feeds did NOT download, so their columns kept last run's numbers:")
        for p in FAILED:
            print("   -", p)
        print("CollegeFootballData.com was slow or down. Run the update again later.")


if __name__ == "__main__":
    main()
