

import os
import re
import sys
import unicodedata


def need(mod):
    try:
        return __import__(mod)
    except ImportError:
        sys.exit(f"Missing package '{mod}'.  Run:  pip install requests openpyxl")


requests = need("requests")
openpyxl = need("openpyxl")


# ---------------------------------------------------------------------------
# Name normalisation
# ---------------------------------------------------------------------------
def norm(s):
    """Lower-case, drop punctuation, expand 'St.'->'state', collapse spaces."""
    s = (s or "").lower().strip()
    # strip accents (José -> jose)
    s = "".join(c for c in unicodedata.normalize("NFKD", s) if not unicodedata.combining(c))
    s = s.replace("&", " and ")
    s = re.sub(r"[.’'()]", "", s)        # . ' ’ ( )
    s = s.replace("-", " ")
    s = re.sub(r"\bst\b", "state", s)
    s = s.replace("university", "").replace("univ", "")
    s = re.sub(r"\s+", " ", s).strip()
    return s


class TeamMatcher:
    """Match a data source's team to the sheet's own team name.

    sheet_names : the names in column A (returned to the client untouched)
    Register every source team with add(); then resolve a sheet name with
    sheet_to_source().
    """

    def __init__(self, sheet_names, alias=None):
        self.sheet_names = list(sheet_names)
        self.alias = alias or {}                      # sheet name -> source abbr/key
        self._by_norm = {}                            # norm(name) -> token
        self._by_abbr = {}                            # ABBR       -> token
        self._disp_tokens = {}                        # token      -> [tokens of long name]
        self._tokens = {}                             # token      -> payload

    def add(self, token, names, abbr=None, payload=None):
        """Register one source team.  `names` = any of its name spellings
        (short, nickname, long).  `abbr` = its abbreviation if it has one."""
        self._tokens[token] = payload
        longest = ""
        for nm in names:
            if not nm:
                continue
            self._by_norm.setdefault(norm(nm), token)
            if len(nm) > len(longest):
                longest = nm
        self._disp_tokens[token] = norm(longest).split()
        if abbr:
            self._by_abbr[abbr.upper()] = token

    def _match(self, sheet_name):
        # 1) explicit alias (by abbreviation or by a spelled-out name)
        if sheet_name in self.alias:
            key = self.alias[sheet_name]
            if key in (None, "", "__NONE__"):
                return None
            if key.upper() in self._by_abbr:
                return self._by_abbr[key.upper()]
            if norm(key) in self._by_norm:
                return self._by_norm[norm(key)]
            return None
        n = norm(sheet_name)
        # 2) exact normalised match
        if n in self._by_norm:
            return self._by_norm[n]
        # 3) the sheet name is itself an abbreviation (JMU, WKU, NIU, ...)
        if sheet_name.upper() in self._by_abbr:
            return self._by_abbr[sheet_name.upper()]
        # 4) prefix-token match (handles a trailing mascot, "Duke Blue Devils")
        nt = n.split()
        hits = [tok for tok, dt in self._disp_tokens.items() if dt[:len(nt)] == nt]
        if len(hits) == 1:
            return hits[0]
        return None

    def sheet_to_source(self):
        """Return (mapping, unmatched) where mapping = {sheet_name: payload}."""
        mapping, unmatched = {}, []
        for name in self.sheet_names:
            tok = self._match(name)
            if tok is None:
                unmatched.append(name)
            else:
                mapping[name] = self._tokens[tok]
        return mapping, unmatched


# ---------------------------------------------------------------------------
# Workbook helpers - only data cells are written; column A is never touched
# ---------------------------------------------------------------------------
def open_sheet(path, sheet_name):
    if not os.path.exists(path):
        sys.exit(f"Can't find '{os.path.basename(path)}' next to the updater.")
    try:                                   # fail early if it's open in Excel
        with open(path, "r+b"):
            pass
    except PermissionError:
        sys.exit(f"'{os.path.basename(path)}' is open in Excel - close it and run again.")
    wb = openpyxl.load_workbook(path)
    return wb, wb[sheet_name]


def sheet_rows(ws, first_data_row=2):
    """{team name in column A: row number} for every non-empty team row."""
    out = {}
    for r in range(first_data_row, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if v not in (None, ""):
            out[v] = r
    return out


def write_values(ws, rows_by_team, col, values_by_team, nd=None):
    """Write one column.  values_by_team maps sheet-team -> value.
    Returns how many cells were filled."""
    n = 0
    for team, row in rows_by_team.items():
        if team in values_by_team:
            v = values_by_team[team]
            if v is None:
                continue
            if nd is not None and isinstance(v, (int, float)):
                v = round(float(v), nd)
            ws.cell(row, col, v)
            n += 1
    return n


def save_sheet(wb, path):
    from openpyxl.workbook.properties import CalcProperties
    wb.calculation = CalcProperties(fullCalcOnLoad=True, calcMode="auto")
    try:
        wb.save(path)
    except PermissionError:
        sys.exit(f"Couldn't save '{os.path.basename(path)}' - is it open in Excel?")


def report(label, unmatched):
    if unmatched:
        print(f"\n  NOTE - {len(unmatched)} team(s) on the sheet got no {label} data")
        print("  (name not found in that source, or it has no number yet):")
        for t in unmatched:
            print("     -", t)


# ---------------------------------------------------------------------------
# Barttorvik - get past the one-time "verifying browser" JS challenge
# ---------------------------------------------------------------------------
_UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
       "(KHTML, like Gecko) Chrome/129.0 Safari/537.36")


def barttorvik_get(url):
    """Return the HTML of a Barttorvik page, solving its JS cookie check."""
    s = requests.Session()
    s.headers.update({"User-Agent": _UA})
    s.get(url, timeout=60)                                   # sets the challenge
    s.post(url, data={"js_test_submitted": "1"}, timeout=60)  # answers it
    r = s.get(url, timeout=60)                               # now the real page
    r.raise_for_status()
    return r.text


def http_get(url):
    """Plain GET with a browser user-agent; returns page text."""
    r = requests.get(url, headers={"User-Agent": _UA}, timeout=60)
    r.raise_for_status()
    return r.text


from html.parser import HTMLParser


class _TableParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.tables, self._rows, self._row, self._cell = [], None, None, None

    def handle_starttag(self, tag, attrs):
        if tag == "table":
            self._rows = []
        elif tag == "tr" and self._rows is not None:
            self._row = []
        elif tag in ("td", "th") and self._row is not None:
            self._cell = []

    def handle_data(self, data):
        if self._cell is not None:
            self._cell.append(data)

    def handle_endtag(self, tag):
        if tag in ("td", "th") and self._cell is not None:
            self._row.append("".join(self._cell).strip())
            self._cell = None
        elif tag == "tr" and self._row is not None:
            if self._row:
                self._rows.append(self._row)
            self._row = None
        elif tag == "table" and self._rows is not None:
            self.tables.append(self._rows)
            self._rows = None


def parse_tables(html):
    """Return every HTML <table> as a list of rows, each row a list of cell texts."""
    p = _TableParser()
    p.feed(html)
    return p.tables
